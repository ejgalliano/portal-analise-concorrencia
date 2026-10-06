"""
Monta o .docx final a partir dos dados já pesquisados e sintetizados.
Este módulo só RENDERIZA — todo o texto já vem pronto do generator.py.
Mantém a mesma ordem de seções validada manualmente (ver report_frentes.ORDEM_SECOES).
"""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .column_labels import rotulo_coluna
from .config import settings
from .report_frentes import FRENTES


def _titulo(doc: Document, texto: str, nivel: int = 1):
    doc.add_heading(texto, level=nivel)


def _paragrafo(doc: Document, texto: str):
    p = doc.add_paragraph(texto)
    p.paragraph_format.space_after = Pt(8)
    return p


def _tabela_simples(doc: Document, cabecalho: list[str], linhas: list[list[str]]):
    tabela = doc.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Light Grid Accent 1"
    hdr = tabela.rows[0].cells
    for i, titulo_coluna in enumerate(cabecalho):
        hdr[i].text = titulo_coluna
    for linha in linhas:
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = str(valor)
    doc.add_paragraph()  # espaçamento depois da tabela


def montar_docx(
    cliente_nome: str,
    resultados: dict[str, dict],
    secoes_gerais: dict,
    caminho_saida: str,
) -> str:
    doc = Document()

    # Capa / título
    titulo = doc.add_heading(
        f"Relatório de Análise de Concorrência — {cliente_nome}", level=0
    )
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"Gerado por {settings.nome_produto} · {settings.nome_agencia} · "
        f"{datetime.now().strftime('%d/%m/%Y')}"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Sumário executivo
    _titulo(doc, "Sumário executivo")
    _paragrafo(doc, secoes_gerais.get("sumario_executivo", ""))

    # Metodologia e fontes
    _titulo(doc, "Metodologia e fontes")
    _paragrafo(doc, secoes_gerais.get("metodologia_e_fontes", ""))

    # Diagnóstico completo por empresa
    _titulo(doc, "Diagnóstico completo por empresa")
    _paragrafo(
        doc,
        "Cada empresa abaixo foi avaliada nas mesmas 13 frentes, a partir das "
        "fontes descritas em “Metodologia e fontes”. Cada empresa é um "
        "capítulo próprio, para facilitar a leitura.",
    )
    for nome_empresa, dados in resultados.items():
        rotulo = " (cliente)" if dados.get("papel") == "cliente" else ""
        _titulo(doc, f"{nome_empresa}{rotulo}", nivel=1)
        frentes_dados = dados.get("frentes", {})
        for f in FRENTES:
            chave = str(f["numero"])
            entrada = frentes_dados.get(chave, {})
            _titulo(doc, f"{f['numero']}. {f['titulo']}", nivel=2)
            texto = entrada.get("texto", "Não pesquisado.")
            _paragrafo(doc, texto)

    # Comparativo consolidado
    _titulo(doc, "Comparativo consolidado")
    for bloco in secoes_gerais.get("comparativo_consolidado", []):
        _titulo(doc, bloco.get("frente", ""), nivel=2)
        linhas = [[linha.get("empresa", ""), linha.get("resumo", "")] for linha in bloco.get("linhas", [])]
        _tabela_simples(doc, ["Empresa", "Resumo"], linhas)

    # Painel de reputação e sinais BR
    _titulo(doc, "Painel de reputação e sinais BR")
    painel_rep = secoes_gerais.get("painel_reputacao_sinais_br", [])
    if painel_rep:
        colunas = list(painel_rep[0].get("colunas", {}).keys())
        cabecalho = ["Empresa"] + [rotulo_coluna(c) for c in colunas]
        linhas = [
            [item.get("empresa", "")] + [item.get("colunas", {}).get(c, "") for c in colunas]
            for item in painel_rep
        ]
        _tabela_simples(doc, cabecalho, linhas)

    # Painel de atividade de marketing digital
    _titulo(doc, "Painel de atividade de marketing digital")
    painel_mkt = secoes_gerais.get("painel_atividade_marketing_digital", [])
    if painel_mkt:
        colunas = list(painel_mkt[0].get("colunas", {}).keys())
        cabecalho = ["Empresa"] + [rotulo_coluna(c) for c in colunas]
        linhas = [
            [item.get("empresa", "")] + [item.get("colunas", {}).get(c, "") for c in colunas]
            for item in painel_mkt
        ]
        _tabela_simples(doc, cabecalho, linhas)

    # Mapa de oportunidade comparativo
    _titulo(doc, "Mapa de oportunidade comparativo")
    _paragrafo(
        doc,
        "✓ = confirmado como presente · — = confirmado como ausente · "
        "n/d = não determinável nesta consulta.",
    )
    mapa = secoes_gerais.get("mapa_oportunidade_comparativo", [])
    if mapa:
        nomes_empresas = list(mapa[0].get("empresas", {}).keys())
        cabecalho = ["Canal / frente"] + nomes_empresas
        linhas = [
            [item.get("canal", "")] + [item.get("empresas", {}).get(n, "n/d") for n in nomes_empresas]
            for item in mapa
        ]
        _tabela_simples(doc, cabecalho, linhas)

    # Fontes citadas
    _titulo(doc, "Fontes citadas")
    fontes = secoes_gerais.get("fontes_citadas", [])
    if fontes:
        _tabela_simples(
            doc,
            ["Fonte", "Link", "Data"],
            [[f.get("nome", ""), f.get("url", ""), f.get("data", "")] for f in fontes],
        )

    # Sobre atualizações
    _titulo(doc, "Sobre atualizações")
    _paragrafo(doc, secoes_gerais.get("sobre_atualizacoes", ""))

    Path(caminho_saida).parent.mkdir(parents=True, exist_ok=True)
    doc.save(caminho_saida)
    return caminho_saida
